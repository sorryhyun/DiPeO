import json
import csv
from pathlib import Path
from typing import Any, Dict, List, Optional, Union
from datetime import datetime
import sqlite3
from docx import Document
import aiofiles
from ..exceptions import ValidationError, FileOperationError
from ..utils.base_service import BaseService
from ...config import UPLOAD_DIR, RESULT_DIR, CONVERSATION_LOG_DIR


class UnifiedFileService(BaseService):
    """Unified file service for all file operations."""
    
    def __init__(self, base_dir: Optional[Path] = None):
        super().__init__()
        self.base_dir = base_dir or Path.cwd()
        # Use directories from config.py instead of creating them
        self.upload_dir = UPLOAD_DIR
        self.result_dir = RESULT_DIR
        self.log_dir = CONVERSATION_LOG_DIR
    
    def read(self, 
             path: str, 
             relative_to: str = "base",
             encoding: str = "utf-8") -> Union[str, Dict[str, Any]]:
        """Read file with automatic format detection."""
        file_path = self._resolve_and_validate_path(path, relative_to)
        
        if not file_path.exists():
            raise FileOperationError(f"File not found: {path}")
        
        if file_path.suffix.lower() == '.docx':
            return self._read_docx(file_path)
        elif file_path.suffix.lower() == '.json':
            with open(file_path, 'r', encoding=encoding) as f:
                return json.load(f)
        elif file_path.suffix.lower() == '.csv':
            return self._read_csv(file_path, encoding)
        else:
            return file_path.read_text(encoding=encoding)
    
    async def write(self,
              path: str,
              content: Any,
              relative_to: str = "results",
              format: Optional[str] = None,
              encoding: str = "utf-8") -> str:
        """Write file with automatic format handling."""
        file_path = self._resolve_and_validate_path(path, relative_to, create_parents=True)
        
        if not format:
            suffix = file_path.suffix.lower()
            if suffix == '.json':
                format = 'json'
            elif suffix == '.csv':
                format = 'csv'
            else:
                format = 'text'
        
        # Write based on format
        if format == 'json':
            await self._write_json(file_path, content, encoding)
        elif format == 'csv':
            await self._write_csv(file_path, content, encoding)
        else:
            await self._write_text(file_path, content, encoding)
        
        # Return relative path from base directory
        return str(file_path.relative_to(self.base_dir))
    
    def write_sqlite(self,
                     db_path: str,
                     table_name: str,
                     data: Union[Dict[str, Any], List[Dict[str, Any]]],
                     relative_to: str = "results") -> str:
        """Write data to SQLite database."""
        file_path = self._resolve_and_validate_path(db_path, relative_to, create_parents=True)
        
        if not file_path.suffix:
            file_path = file_path.with_suffix('.db')
        
        # Normalize data to list
        records = [data] if isinstance(data, dict) else data
        if not records:
            raise ValidationError("No data to insert")
        
        # Connect and create table if needed
        conn = sqlite3.connect(str(file_path))
        try:
            columns = list(records[0].keys())
            
            # Create table if it doesn't exist
            col_defs = ", ".join([f'"{col}" TEXT' for col in columns])
            conn.execute(f'CREATE TABLE IF NOT EXISTS "{table_name}" ({col_defs})')
            
            # Insert data
            placeholders = ", ".join(["?" for _ in columns])
            col_names = ", ".join([f'"{col}"' for col in columns])
            
            for record in records:
                values = [str(record.get(col, "")) for col in columns]
                conn.execute(
                    f'INSERT INTO "{table_name}" ({col_names}) VALUES ({placeholders})',
                    values
                )
            
            conn.commit()
            
        finally:
            conn.close()
        
        return str(file_path.relative_to(self.base_dir))
    
    # Removed save_conversation_log - MemoryService handles conversation logging
    
    def list_files(self, 
                   directory: str = ".",
                   relative_to: str = "base",
                   pattern: str = "*") -> List[Dict[str, Any]]:
        """List files in a directory."""
        dir_path = self._resolve_and_validate_path(directory, relative_to)
        
        if not dir_path.is_dir():
            raise ValidationError(f"Not a directory: {directory}")
        
        files = []
        for file_path in dir_path.glob(pattern):
            if file_path.is_file():
                stat = file_path.stat()
                files.append({
                    "name": file_path.name,
                    "path": str(file_path.relative_to(self.base_dir)),
                    "size": stat.st_size,
                    "modified": datetime.fromtimestamp(stat.st_mtime).isoformat()
                })
        
        return sorted(files, key=lambda x: x["modified"], reverse=True)
    
    def delete(self, path: str, relative_to: str = "base") -> bool:
        """Delete a file safely."""
        file_path = self._resolve_and_validate_path(path, relative_to)
        
        if file_path.exists():
            if file_path.is_file():
                file_path.unlink()
                return True
            else:
                raise ValidationError("Cannot delete directories")
        
        return False
    
    def _resolve_and_validate_path(self, 
                                   path: str, 
                                   relative_to: str,
                                   create_parents: bool = False) -> Path:
        """Resolve and validate path security."""
        base_map = {
            "base": self.base_dir,
            "uploads": self.upload_dir,
            "results": self.result_dir,
            "logs": self.log_dir
        }
        
        base_path = base_map.get(relative_to, self.base_dir)
        
        # Resolve the path
        if Path(path).is_absolute():
            resolved = Path(path).resolve()
        else:
            resolved = (base_path / path).resolve()
        
        # Validate it's within allowed directory
        try:
            resolved.relative_to(self.base_dir)
        except ValueError:
            raise ValidationError(f"Path outside allowed directory: {path}")
        
        # Create parent directories if requested
        if create_parents and not resolved.parent.exists():
            resolved.parent.mkdir(parents=True, exist_ok=True)
        
        return resolved
    
    def _read_docx(self, file_path: Path) -> str:
        """Read content from DOCX file."""
        doc = Document(str(file_path))
        return "\n".join([para.text for para in doc.paragraphs])
    
    def _read_csv(self, file_path: Path, encoding: str) -> List[Dict[str, Any]]:
        """Read CSV file and return as list of dictionaries."""
        with open(file_path, 'r', encoding=encoding) as f:
            reader = csv.DictReader(f)
            return list(reader)
    
    async def _write_json(self, file_path: Path, content: Any, encoding: str):
        """Write content as formatted JSON."""
        async with aiofiles.open(file_path, 'w', encoding=encoding) as f:
            await f.write(json.dumps(content, indent=2, ensure_ascii=False))
    
    async def _write_csv(self, file_path: Path, content: Any, encoding: str):
        """Write content as CSV."""
        if isinstance(content, dict):
            rows = [content]
        elif isinstance(content, list) and all(isinstance(r, dict) for r in content):
            rows = content
        else:
            raise ValidationError("CSV content must be a dict or list of dicts")
        
        if not rows:
            raise ValidationError("No data to write to CSV")
        
        # Get all unique keys
        all_keys = set()
        for row in rows:
            all_keys.update(row.keys())
        
        fieldnames = sorted(all_keys)
        
        async with aiofiles.open(file_path, 'w', encoding=encoding, newline='') as f:
            # Write CSV content as string
            import io
            output = io.StringIO()
            writer = csv.DictWriter(output, fieldnames=fieldnames)
            writer.writeheader()
            writer.writerows(rows)
            await f.write(output.getvalue())
    
    async def _write_text(self, file_path: Path, content: Any, encoding: str):
        """Write content as plain text."""
        text_content = str(content) if not isinstance(content, str) else content
        async with aiofiles.open(file_path, 'w', encoding=encoding) as f:
            await f.write(text_content)